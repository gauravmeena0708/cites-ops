"""HTML and Excel views of the shared functionality and issue-nature data."""

import json
import re
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from ..core.top_natures import build_rankings


TEMPLATES = Path(__file__).parents[1] / "templates"


def write_html(data, output, template="briefing.html"):
    # JSON in a script element must not allow source text to terminate that element.
    payload = json.dumps(data, ensure_ascii=False, allow_nan=False)
    payload = payload.replace("<", "\\u003c").replace(">", "\\u003e").replace("&", "\\u0026")
    payload = payload.replace("\u2028", "\\u2028").replace("\u2029", "\\u2029")
    Path(output).write_text((TEMPLATES / template).read_text(encoding="utf-8").replace("__PAYLOAD__", payload), encoding="utf-8")


def write_workbook(data, output):
    wb = Workbook()
    wb.remove(wb.active)

    def sheet(name, headers, rows, widths, note):
        ws = wb.create_sheet(name)
        ws.append([f"CITES — {name} — {data['date']}"])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        ws.cell(1, 1).font = Font(name="Calibri", size=16, bold=True, color="17354D")
        ws.append([note])
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
        ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[2].height = 46
        ws.append(headers)
        for cell in ws[3]:
            cell.fill = PatternFill("solid", fgColor="17354D")
            cell.font = Font(name="Calibri", color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True)
        ws.row_dimensions[3].height = 30
        for row_index, values in enumerate(rows, 4):
            ws.append(values)
            if name != "Issue Details":
                ws.row_dimensions[row_index].height = 34
            for cell in next(ws.iter_rows(min_row=row_index, max_row=row_index, max_col=len(headers))):
                if isinstance(cell.value, str):
                    # Treat tracker/chat strings as text, including Excel formula prefixes.
                    cell.value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", cell.value)[:32767]
                    cell.data_type = "s"
                if isinstance(cell.value, date):
                    cell.number_format = "dd-mmm-yyyy"
                cell.font = Font(name="Calibri", size=11)
                cell.alignment = Alignment(vertical="top", wrap_text=name != "Issue Details")
                if row_index % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F0F5F8")
        for index, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(index)].width = width
        ws.freeze_panes = "B4"
        ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{max(3, ws.max_row)}"
        ws.sheet_view.showGridLines = False
        ws.print_title_rows = "1:3"
        return ws

    scope = "All exported tickets" if data["scope"] == "all" else "Currently open tickets only" if data["scope"] == "open" else "Filtered tickets only"
    period = f"Last week: {data['week']['start']} to {data['week']['end']} (submission date). {scope}."
    count_note = f"Counts: {data['count_source']}. Blank counts are unavailable, not zero. " + " ".join(data["warnings"])
    totals = ["OVERALL"] + [data["totals"][k] for k in ("total", "open", "resolved", "closed")] + ["", "", "", len(data["issues"])]
    sheet("Functionality Summary", ["Functionality", "Total", "Open", "Resolved", "Closed", "Responsible Officer", "DD", "JD", "Tickets in CSV"],
          [totals] + [[r[k] for k in ("functionality", "total", "open", "resolved", "closed", "officer", "dd", "jd", "detail_total")] for r in data["functions"]],
          [34, 13, 13, 13, 13, 32, 28, 28, 17], count_note)
    nature_rows = sorted(data["natures"], key=lambda r: (r["functionality"], r.get("rank", 9999), r["nature"]))
    sheet("Issue Natures", ["Functionality", "Issue Nature", "Total", "Open", "Resolved", "Closed", "Last Week", "Previous Week", "Change", "Top 5 Backlog", "Top 5 Last Week"],
          [[r["functionality"], r["nature"], r["total"], r["open"], r["resolved"], r["closed"], r["last_week"], r["prior_week"],
            r["last_week"] - r["prior_week"], "Yes" if r.get("rank", 99) <= 5 else "", "Yes" if r.get("week_rank", 99) <= 5 else ""] for r in nature_rows],
          [30, 65, 12, 12, 12, 12, 14, 16, 12, 18, 18], period + " Counts describe current status; weekly change is intake, not resolution events. All natures are included for reconciliation.")
    cols = ("id", "functionality", "summary", "status", "assigned_to", "officer", "nature", "submitted", "updated", "age_days", "dd", "jd", "tracker_status", "description")
    rows = []
    for issue in data["issues"]:
        rows.append([date.fromisoformat(issue[k]) if k in ("submitted", "updated") and issue[k] else issue[k] for k in cols])
    ws = sheet("Issue Details", ["ID", "Functionality", "Summary", "Status", "Assigned Queue", "Responsible Officer", "Issue Nature", "Submitted", "Updated", "Age (Days)", "DD", "JD", "Tracker Status", "Description"],
               rows, [13, 28, 65, 14, 30, 30, 55, 16, 16, 14, 28, 28, 18, 90], period + " Expand hidden columns K:N for management hierarchy, original status and full description.")
    for col in ("K", "L", "M", "N"):
        ws.column_dimensions[col].hidden = True
    ranking = data.get("rankings") or build_rankings(data)
    ranking_period = "Full supplied snapshot" if ranking["period"] == "all" else period
    for key, title in (("open", "Top Open Natures"), ("resolved_closed", "Top Resolved Closed")):
        group = ranking[key]
        note = f"Top {ranking['top_n']} issue natures. {ranking_period}. {scope}. Distinct tickets; review bucket excluded from ranking: {group['review_count']}."
        if not group["available"]:
            note = "Unavailable: the supplied CSV contains only open tickets. Supply completed-ticket details for this snapshot."
        output_rows = []
        for row in group["rows"]:
            samples = [f"#{s['id']} | {s['functionality']} | {s['status']} | {s['summary']}" for s in row["samples"]]
            samples += [""] * (3 - len(samples))
            functions = "; ".join(f"{f['name']} ({f['count']})" for f in row["functionalities"])
            output_rows.append([row["rank"], row["nature"], row["count"], row["resolved"], row["closed"], functions, *samples])
        ws = sheet(title, ["Rank", "Issue Nature", "Ticket Count", "Resolved", "Closed", "Affected Functionalities", "Example 1", "Example 2", "Example 3"],
                   output_rows, [9, 58, 15, 13, 13, 50, 65, 65, 65], note)
        for row_index in range(4, 4 + len(output_rows)):
            ws.row_dimensions[row_index].height = 92
        if not group["rows"]:
            ws.append(["No ranked issue natures available for this scope."])
    wb.save(output)
