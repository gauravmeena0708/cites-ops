import json
import tempfile
import unittest
import zipfile
from datetime import date
from pathlib import Path
from unittest.mock import patch

import pandas as pd
import yaml
from openpyxl import load_workbook

from functools import partial
from cites_ops.brief import run_brief as full_run_brief
run_brief = partial(full_run_brief, no_pptx=True)
from cites_ops.core.briefing import prepare_briefing
from cites_ops.core.chat_parser import ChatParser
from cites_ops.core.issue_nature import IssueNatureClassifier, OTHER
from cites_ops.core.knowledge import extract_candidates
from cites_ops.core.top_natures import build_rankings
from cites_ops.reporters.briefing_reporter import write_html, write_workbook


DATE = date(2026, 9, 7)


def ticket(identifier="1", status="assigned", submitted="31-08-2026", summary="Claim not visible at DA level"):
    return {"Id": identifier, "Project": "CITES", "Category": "Form-13", "Status": status,
            "Summary": summary, "Description": "Claim missing at dealing assistant login",
            "Date Submitted": submitted, "Updated": "07-09-2026", "Assigned To": "team_epfo_form_13"}


def teams():
    return pd.DataFrame([{"Team": "Form-13", "Account handled by": "Officer A", "DD(IS)": "Deputy A", "JD(IS)": "Joint A"}])


def official():
    metrics = {"total": 4, "open": 2, "resolved": 1, "closed": 1}
    return {"source": {"name": "status.docx", "data_date": "2026-09-07"}, "totals": metrics,
            "categories": [{"module_label": "Form-13", **metrics}]}


class TestBriefData(unittest.TestCase):
    def prepare(self, rows, **kwargs):
        return prepare_briefing(pd.DataFrame(rows), teams(), IssueNatureClassifier(), DATE, **kwargs)

    def test_four_status_counts_and_calendar_week_boundaries(self):
        data = self.prepare([ticket(), ticket("2", " fixed ", "06-09-2026"),
                             ticket("3", "closed", "30-08-2026"), ticket("4", "new", "07-09-2026")], scope="all")
        self.assertEqual(data["totals"], {"total": 4, "open": 2, "resolved": 1, "closed": 1})
        self.assertEqual(data["week"]["start"], "2026-08-31")
        self.assertEqual(data["week"]["end"], "2026-09-06")
        self.assertEqual(sum(i["last_week"] for i in data["issues"]), 2)
        self.assertEqual(sum(i["prior_week"] for i in data["issues"]), 1)
        self.assertEqual(sum(r["total"] for r in data["natures"]), 4)
        self.assertEqual(data["issues"][0]["officer"], "Officer A")
        self.assertEqual(data["issues"][0]["age_days"], 7)

    def test_open_csv_uses_reconciled_portfolio_counts_without_invented_detail(self):
        data = self.prepare([ticket(), ticket("2", "feedback")], stats=official())
        self.assertEqual(data["scope"], "open")
        self.assertEqual(data["totals"]["total"], 4)
        self.assertEqual(data["detail_counts"]["total"], 2)
        self.assertEqual(data["functions"][0]["resolved"], 1)
        self.assertEqual(sum(r["resolved"] for r in data["natures"]), 0)

    def test_open_without_stats_leaves_unknown_counts_unavailable(self):
        data = self.prepare([ticket()])
        self.assertEqual(data["totals"], {"total": None, "open": 1, "resolved": None, "closed": None})
        self.assertIsNone(data["functions"][0]["total"])

    def test_mixed_status_auto_requires_scope_and_partial_never_claims_portfolio(self):
        with self.assertRaisesRegex(ValueError, "--scope"):
            self.prepare([ticket(), ticket("2", "closed")])
        data = self.prepare([ticket(), ticket("2", "closed")], scope="partial")
        self.assertTrue(all(v is None for v in data["totals"].values()))

    def test_invalid_or_inconsistent_sources_are_rejected(self):
        with self.assertRaisesRegex(ValueError, "disagree"):
            self.prepare([ticket()], stats=official())
        bad = official()
        bad["source"]["data_date"] = "2026-09-06"
        with self.assertRaisesRegex(ValueError, "date"):
            self.prepare([ticket(), ticket("2")], stats=bad)
        with self.assertRaisesRegex(ValueError, "Unrecognized"):
            self.prepare([ticket(status="mystery")])
        with self.assertRaisesRegex(ValueError, "submitted after"):
            self.prepare([ticket(submitted="08-09-2026")])
        row = ticket()
        row["Updated"] = "08-09-2026"
        with self.assertRaisesRegex(ValueError, "updated after"):
            self.prepare([row])

    def test_identical_duplicate_dedup_but_conflicting_ids_fail(self):
        data = self.prepare([ticket(), ticket()])
        self.assertEqual(len(data["issues"]), 1)
        with self.assertRaisesRegex(ValueError, "duplicate"):
            self.prepare([ticket(), ticket("0001", summary="Different")])

    def test_invalid_dates_remain_unknown_and_do_not_enter_week(self):
        data = self.prepare([ticket(submitted="unknown")])
        self.assertIsNone(data["issues"][0]["age_days"])
        self.assertFalse(data["issues"][0]["last_week"])
        self.assertTrue(any("submission dates" in warning for warning in data["warnings"]))

    def test_conflicting_ownership_fails_instead_of_selecting_first(self):
        owners = pd.concat([teams(), teams().assign(**{"Account handled by": "Someone else"})])
        with self.assertRaisesRegex(ValueError, "Conflicting ownership"):
            prepare_briefing(pd.DataFrame([ticket()]), owners, IssueNatureClassifier(), DATE)

    def test_review_bucket_does_not_masquerade_as_ranked_nature(self):
        row = ticket(summary="Please help")
        row["Description"] = ""
        data = self.prepare([row])
        self.assertEqual(data["issues"][0]["nature"], OTHER)
        self.assertNotIn("rank", data["natures"][0])


class TestNatureRules(unittest.TestCase):
    def test_scoping_exclusions_and_cache_invalidation(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "rules.yaml"
            config = {"version": "test", "rules": [{"rule_id": "R1", "label": "Missing claim",
                "priority": 100, "functionalities": ["Form-13"], "patterns": ["claim.*visible"],
                "exclude_patterns": ["now visible"]}]}
            path.write_text(yaml.safe_dump(config), encoding="utf-8")
            classifier = IssueNatureClassifier(path)
            self.assertEqual(classifier.classify("Form 13", "claim not visible", "")["nature"], "Missing claim")
            self.assertEqual(classifier.classify("Form-13", "claim now visible", "")["nature"], OTHER)
            self.assertEqual(classifier.classify("Form-19", "claim not visible", "")["nature"], OTHER)
            classifier.classify("Form 13", "claim not visible", "")
            self.assertEqual(classifier.hits, 1)
            config["rules"][0]["label"] = "Revised label"
            path.write_text(yaml.safe_dump(config), encoding="utf-8")
            changed = IssueNatureClassifier(path, classifier.cache)
            self.assertEqual(changed.classify("Form 13", "claim not visible", "")["nature"], "Revised label")
            self.assertEqual(changed.hits, 0)


class TestBriefOutputs(unittest.TestCase):
    def test_excel_types_minimal_sheets_and_safe_html_payload(self):
        row = ticket(summary="=HYPERLINK(\"https://invalid\")")
        row["Description"] = "</script><script>alert('test')</script>"
        data = prepare_briefing(pd.DataFrame([row]), teams(), IssueNatureClassifier(), DATE, "all")
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            write_workbook(data, root / "issues.xlsx")
            wb = load_workbook(root / "issues.xlsx")
            self.assertEqual(wb.sheetnames, ["Functionality Summary", "Issue Natures", "Issue Details", "Top Open Natures", "Top Resolved Closed"])
            ws = wb["Issue Details"]
            self.assertEqual(ws["C4"].data_type, "s")
            self.assertTrue(ws["C4"].value.startswith("=HYPERLINK"))
            self.assertEqual(ws["J4"].value, 7)
            self.assertEqual(ws["J4"].data_type, "n")
            self.assertEqual(ws["H4"].value.date(), date(2026, 8, 31))
            self.assertTrue(ws.column_dimensions["N"].hidden)
            wb.close()
            write_html(data, root / "dashboard.html")
            html = (root / "dashboard.html").read_text(encoding="utf-8")
            self.assertNotIn("</script><script>alert", html)
            payload = html.split('<script id="report-data" type="application/json">')[1].split('</script>')[0]
            self.assertEqual(json.loads(payload)["issues"][0]["description"], row["Description"])

    @patch("urllib.request.urlopen", side_effect=AssertionError("Offline workflow must not access network"))
    def test_folder_workflow_two_outputs_cache_and_safe_overwrite(self, network):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            inputs = root / "input"
            inputs.mkdir()
            pd.DataFrame([ticket()]).to_csv(inputs / "issues.csv", index=False)
            teams().to_csv(inputs / "teams.csv", index=False)
            final, manifest = run_brief(inputs, root / "out", DATE)
            self.assertEqual({p.name for p in final.iterdir() if p.is_file()}, {"dashboard_2026-09-07.html", "issues_2026-09-07.xlsx"})
            self.assertEqual(manifest["cache_misses"], 1)
            with self.assertRaises(FileExistsError):
                run_brief(inputs, root / "out", DATE)
            _, second = run_brief(inputs, root / "out", DATE, overwrite=True)
            self.assertEqual(second["cache_hits"], 1)
            self.assertFalse(any((root / "out").glob(".brief-*")))
            self.assertFalse(network.called)

    def test_history_uses_actual_prior_date_and_does_not_compare_open_only_totals(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            inputs = root / "input"
            inputs.mkdir()
            rows = [ticket()]
            rows[0]["Updated"] = "03-09-2026"
            pd.DataFrame(rows).to_csv(inputs / "issues.csv", index=False)
            teams().to_csv(inputs / "teams.csv", index=False)
            run_brief(inputs, root / "out", date(2026, 9, 3), scope="all")
            rows.append(ticket("2", "closed"))
            pd.DataFrame(rows).to_csv(inputs / "issues.csv", index=False)
            final, _ = run_brief(inputs, root / "out", DATE, scope="all")
            html = (final / "dashboard_2026-09-07.html").read_text(encoding="utf-8")
            data = json.loads(html.split('<script id="report-data" type="application/json">')[1].split('</script>')[0])
            self.assertEqual(data["comparison"]["date"], "2026-09-03")
            self.assertEqual(data["comparison"]["delta"]["closed"], 1)
            pd.DataFrame([ticket()]).to_csv(inputs / "issues.csv", index=False)
            final, _ = run_brief(inputs, root / "out", DATE, scope="open", overwrite=True)
            html = (final / "dashboard_2026-09-07.html").read_text(encoding="utf-8")
            data = json.loads(html.split('<script id="report-data" type="application/json">')[1].split('</script>')[0])
            self.assertIsNone(data["comparison"])

    def test_output_cannot_replace_input_or_unrelated_directory(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            inputs = root / DATE.isoformat()
            inputs.mkdir()
            pd.DataFrame([ticket()]).to_csv(inputs / "issues.csv", index=False)
            teams().to_csv(inputs / "teams.csv", index=False)
            with self.assertRaisesRegex(ValueError, "input files"):
                run_brief(inputs, root, DATE, overwrite=True)
            other = root / "out" / DATE.isoformat()
            other.mkdir(parents=True)
            (other / "keep.txt").write_text("User file", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "not a simplified briefing"):
                run_brief(inputs, root / "out", DATE, overwrite=True)
            self.assertEqual((other / "keep.txt").read_text(encoding="utf-8"), "User file")
            self.assertTrue((inputs / "issues.csv").is_file())


class TestTopNatures(unittest.TestCase):
    def test_status_separation_review_counts_and_weekly_submission_scope(self):
        rows = [ticket(), ticket("2", "resolved"), ticket("3", "closed", "30-08-2026"),
                ticket("4", summary="No recognizable symptom")]
        rows[-1]["Description"] = ""
        data = prepare_briefing(pd.DataFrame(rows), teams(), IssueNatureClassifier(), DATE, "all")
        ranked = build_rankings(data, 10)
        self.assertEqual(ranked["open"]["total"], 2)
        self.assertEqual(ranked["open"]["review_count"], 1)
        self.assertEqual(ranked["open"]["ranked_total"], 1)
        completed = ranked["resolved_closed"]
        self.assertEqual(completed["total"], 2)
        self.assertEqual(completed["rows"][0]["count"], 2)
        self.assertEqual(completed["rows"][0]["resolved"], 1)
        self.assertEqual(completed["rows"][0]["closed"], 1)
        weekly = build_rankings(data, 15, "week")
        self.assertEqual(weekly["resolved_closed"]["total"], 1)
        # A recent update never moves an old submission into the weekly ranking.
        self.assertEqual(weekly["resolved_closed"]["closed"], 0)
        open_data = prepare_briefing(pd.DataFrame([ticket()]), teams(), IssueNatureClassifier(), DATE)
        self.assertFalse(build_rankings(open_data)["resolved_closed"]["available"])

    def test_top_limit_ties_distinct_examples_and_masking(self):
        rows = []
        for index in range(18):
            row = ticket(str(index+1), summary=f"Claim not visible, example {index}; email a@example.org; UAN 100252164386")
            row["Description"] = "Claim missing. PW:SampleSecret123; User name: sampleuser"
            rows.append(row)
        data = prepare_briefing(pd.DataFrame(rows), teams(), IssueNatureClassifier(), DATE, "all")
        for index, row in enumerate(data["issues"]):
            row["nature"] = f"Nature {index:02}"
        ranked = build_rankings(data, 10)
        self.assertEqual(len(ranked["open"]["rows"]), 10)
        self.assertEqual(ranked["open"]["rows"][0]["nature"], "Nature 00")
        self.assertEqual(len(build_rankings(data, 15)["open"]["rows"]), 15)
        # Identical source text is not repeated to fill three example slots.
        for row in data["issues"]:
            row["nature"] = "Visibility"
            row["summary"] = data["issues"][0]["summary"]
        examples = build_rankings(data)["open"]["rows"][0]["samples"]
        self.assertEqual(len(examples), 1)
        self.assertNotIn("a@example.org", examples[0]["summary"])
        self.assertNotIn("100252164386", examples[0]["summary"])
        self.assertNotIn("SampleSecret123", examples[0]["description"])
        self.assertNotIn("sampleuser", examples[0]["description"])

    def test_excel_ranked_counts_match_shared_data(self):
        data = prepare_briefing(pd.DataFrame([ticket(), ticket("2", "closed")]), teams(), IssueNatureClassifier(), DATE, "all")
        data["rankings"] = build_rankings(data, 10)
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "rankings.xlsx"
            write_workbook(data, output)
            wb = load_workbook(output)
            for key, sheet in (("open", "Top Open Natures"), ("resolved_closed", "Top Resolved Closed")):
                expected = data["rankings"][key]["rows"][0]
                self.assertEqual(wb[sheet]["B4"].value, expected["nature"])
                self.assertEqual(wb[sheet]["C4"].value, expected["count"])
                self.assertIn(expected["samples"][0]["id"], wb[sheet]["G4"].value)
            wb.close()

    @patch("cites_ops.reporters.top_natures_pptx.write_presentations", side_effect=RuntimeError("Renderer failed"))
    def test_ppt_failure_preserves_previous_report(self, renderer):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            inputs = root / "input"
            inputs.mkdir()
            pd.DataFrame([ticket()]).to_csv(inputs / "issues.csv", index=False)
            teams().to_csv(inputs / "teams.csv", index=False)
            final, _ = run_brief(inputs, root / "out", DATE)
            before = (final / "dashboard_2026-09-07.html").read_bytes()
            with self.assertRaisesRegex(RuntimeError, "Renderer failed"):
                full_run_brief(inputs, root / "out", DATE, overwrite=True)
            self.assertEqual((final / "dashboard_2026-09-07.html").read_bytes(), before)
            self.assertFalse(any((root / "out").glob(".brief-*")))


class TestKnowledgeCandidates(unittest.TestCase):
    def test_announcements_not_vague_replies_negations_or_promises(self):
        data = prepare_briefing(pd.DataFrame([ticket()]), teams(), IssueNatureClassifier(), DATE)
        messages = [
            "31/08/2026, 10:00 - A: Form-13 claim issue reported.",
            "31/08/2026, 10:05 - B: Form-13 issue resolved. Ticket 0001. Contact person@example.com or 9876543210",
            "31/08/2026, 10:06 - B: done",
            "31/08/2026, 10:07 - B: Form-13 issue not resolved",
            "31/08/2026, 10:08 - B: Form-13 issue will be resolved",
            "31/08/2026, 10:09 - B: Is the Form-13 issue resolved?",
            "31/08/2026, 10:10 - B: New functionality deployed for Form-13.",
            "31/08/2026, 10:11 - B: Form-13 workaround: follow these steps to retry.",
            "08/09/2026, 10:11 - B: Form-13 issue resolved.",
        ]
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "chat.txt"
            path.write_text("\n".join(messages), encoding="utf-8")
            entries, warnings = extract_candidates([path], data["issues"], date(2026, 8, 31), DATE)
            self.assertEqual(len(entries), 3)
            self.assertEqual({e["type"] for e in entries}, {"New functionality", "Reported fix", "Workaround"})
            self.assertTrue(all("review" in e["state"] for e in entries))
            fix = next(e for e in entries if e["type"] == "Reported fix")
            self.assertEqual(fix["tickets"], ["1"])
            self.assertGreater(len(fix["evidence"]), 1)
            self.assertNotIn("9876543210", json.dumps(entries))
            self.assertNotIn("person@example.com", json.dumps(entries))
            self.assertEqual(warnings, [])

    def test_zip_reads_all_transcripts_and_system_events_do_not_merge(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "chats.zip"
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("one.txt", "31/08/2026, 10:00 - A: First message\n31/08/2026, 10:01 - B joined using a link\n31/08/2026, 10:02 - A: Next message")
                archive.writestr("two.txt", "31/08/2026, 11:00 - C: Another transcript")
            df = ChatParser.parse_file(path)
            self.assertEqual(len(df), 3)
            self.assertEqual(set(df["source_file"]), {"one.txt", "two.txt"})
            self.assertNotIn("joined", " ".join(df["message"]))


if __name__ == "__main__":
    unittest.main()
